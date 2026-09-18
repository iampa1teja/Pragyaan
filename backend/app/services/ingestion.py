import base64
from pathlib import Path
from dataclasses import dataclass
from pypdf import PdfReader
from docx import Document

from .embeddings import embed_texts
from .llm import get_client
from ..core.config import get_settings
from ..store.chroma_store import ChromaStore
from ..store.mongo import stories

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"}
_MIME = {".png": "png", ".jpg": "jpeg", ".jpeg": "jpeg", ".webp": "webp", ".gif": "gif", ".bmp": "bmp"}

#--------Parser--------
def parse_text(path: Path) -> str:
    return path.read_text(encoding = "utf-8")

def parse_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    pages = []

    for page in reader.pages:
        text = page.extract_text()
        if text:
            pages.append(text)
    return  "\n\n".join(pages)

def parse_docx(path: Path) -> str:
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)

async def parse_image(path: Path) -> str:
    """Transcribe/describe an image using the Ollama vision model."""
    mime = _MIME.get(path.suffix.lower(), "png")
    b64 = base64.b64encode(path.read_bytes()).decode("ascii")
    resp = await get_client().chat.completions.create(
        model=get_settings().vision_model,
        messages=[{
            "role": "user",
            "content": [
                {"type": "text", "text": "Transcribe all text in this image exactly. If there is little or no text, describe the scene in detail so it can be used as story context."},
                {"type": "image_url", "image_url": {"url": f"data:image/{mime};base64,{b64}"}},
            ],
        }],
    )
    return resp.choices[0].message.content or ""

def parse_file(path: Path) -> str:
    suffix = path.suffix.lower()

    if suffix in {".txt", ".md"}:
        return parse_text(path)
    if suffix == ".pdf":
        return parse_pdf(path)
    if suffix == ".docx":
        return parse_docx(path)
    raise ValueError(f"Unsupported file type: {suffix}")
#-----------------------

#-------Chunking--------
@dataclass
class Chunk:
    id: str
    text: str
    char_start: int
    char_end: int

def chunk_text(text: str, size: int = 1000, overlap: int = 150, ) -> list[Chunk]:
    if size <= 0:
        raise ValueError("size must be greater than 0")

    if overlap < 0:
        raise ValueError("overlap cannot be negative")

    if overlap >= size:
        raise ValueError("overlap must be smaller than size")

    chunks = []
    step = size - overlap
    for start in range(0, len(text), step):
        end = min(start + size, len(text))
        chunk = Chunk(
            id=f"chunk_{len(chunks)}",
            text=text[start:end],
            char_start=start,
            char_end=end,
        )
        chunks.append(chunk)
        if end == len(text):
            break
    return chunks

#-----------------------


async def _extract_one(path: Path) -> str:
    if path.suffix.lower() in IMAGE_EXTS:
        return await parse_image(path)
    return parse_file(path)


async def ingest_story(story_id: str, file_paths: Path | list[Path]) -> None:
    if isinstance(file_paths, Path):
        file_paths = [file_paths]

    parts = []
    for p in file_paths:
        t = await _extract_one(p)
        if t.strip():
            parts.append(t.strip())
    text = "\n\n".join(parts)

    if not text.strip():
        raise ValueError("Files contain no extractable text")
    chunks = chunk_text(text)
    if not chunks:
        raise ValueError("No chunks generated from file")

    documents = [chunk.text for chunk in chunks]
    embeddings = await embed_texts(documents)

    ids = [chunk.id for chunk in chunks]

    metadatas = [
        {
            "story_id": story_id,
            "char_start": chunk.char_start,
            "char_end": chunk.char_end,
        }
        for chunk in chunks
    ]

    store = ChromaStore()
    store.add_chunks(
        story_id=story_id,
        ids=ids,
        docs=documents,
        metadatas=metadatas,
        embeddings=embeddings,
    )

    await stories().update_one(
        {"_id": story_id},
        {"$set": {"status": "ingested", "n_chunks": len(chunks)}},
    )